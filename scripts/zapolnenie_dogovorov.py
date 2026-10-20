#!/usr/bin/env python3
"""Автозаполнение трудовых договоров и приказов (.docx) из Excel.

Запуск:   python3 zapolnenie_dogovorov.py
Нужно:    pip install python-docx openpyxl

Как работает
  * Спрашивает: какой Excel-файл, какой ЛИСТ, какие шаблоны .docx (договор, приказ ...),
    в какой СТРОКЕ шапка (заголовки) и в КАКОЙ КОЛОНКЕ лежит каждый ключ ({ФИО}, {РМА} ...).
    Совпадающие по названию колонки подставляются сами — остаётся нажать Enter.
    Выбор сохраняется в mapping_<файл>_<лист>.json и предлагается при следующем запуске.
  * Каждая СТРОКА под шапкой = один человек. Для каждого человека создаётся своя папка.
  * Берутся ТОЛЬКО действующие работники: у кого колонка «Санаи хориҷ» (дата увольнения) ПУСТАЯ.
    Колонка подбирается сама (Санаи хориҷ / Санаи хорич / Дата увольнения) или выбирается вместе с остальными
    ключами; уволенные пропускаются и перечислены в отчёте. Чтобы взять всех — запуск с флагом --all.
  * Пустая ячейка (или нет такой колонки) -> в документе остаются пробелы.
  * ФИО: везде, где код начинается с «ФИО», выводится «Кириллица (Pinyin)».
        - есть кириллица и латиница в колонках ФИО и ФИО_ЛАТ -> берутся обе;
        - есть только латиница (пиньин) -> кириллица создаётся автотранслитом (в отчёте помечено «проверить»);
        - есть только кириллица -> выводится она одна.
  * Даты: САНА / САНА_СОЛИ / СОЛ / МОҲ / РӮЗ можно брать из одной колонки с датой.
  * В конце — отчёт по каждому человеку и каждому шаблону: что заполнено, что пусто, чего нет в Excel.
    Терминал не засоряется: пишутся только короткие итоги. Подробный отчёт — report.html (красивый, открывается
    в браузере) и report.xlsx (таблица с цветами) в папке результата.
"""
import argparse
import datetime as dt
import html
import json
import os
import re
import sys
import webbrowser

import openpyxl
from docx import Document

SPACES = " " * 14                       # чем заменять пустые значения
FIO_BOTH_EVERYWHERE = True              # во всех кодах ФИО* писать «Кирилица (Pinyin)»
PH = re.compile(r"\{([^{}]+)\}")

MONTHS_TJ = ["январи", "феврали", "марти", "апрели", "майи", "июни",
             "июли", "августи", "сентябри", "октябри", "ноябри", "декабри"]

# Дополнительные названия колонок для одного и того же кода
ALIASES = {
    "岗位": ["岗位", "ВАЗИФА_КИТ", "ВАЗИФА_ЗХ"],
}

# ----------------------------------------------------------------- транслит
_LAT2TJ_DI = {"zh": "ж", "ch": "ч", "sh": "ш", "ng": "нг", "yu": "ю", "ya": "я",
              "yo": "ё", "ye": "е", "xi": "си", "qi": "чи", "ji": "ҷи", "ai": "ай",
              "ei": "эй", "ao": "ао", "ou": "оу", "ui": "уй", "iu": "ю"}
_LAT2TJ = dict(a="а", b="б", c="ц", d="д", e="е", f="ф", g="г", h="ҳ", i="и",
               j="ҷ", k="к", l="л", m="м", n="н", o="о", p="п", q="ч", r="р",
               s="с", t="т", u="у", v="в", w="в", x="кс", y="й", z="з")


def latin_to_tj(s: str) -> str:
    """Грубый транслит латиницы (пиньин) в таджикскую кириллицу. Всегда проверять вручную."""
    out = []
    for word in s.split():
        w, i, res = word.lower(), 0, ""
        while i < len(w):
            if w[i:i + 2] in _LAT2TJ_DI:
                res += _LAT2TJ_DI[w[i:i + 2]]
                i += 2
            else:
                res += _LAT2TJ.get(w[i], w[i])
                i += 1
        out.append(res[:1].upper() + res[1:])
    return " ".join(out)


def title_pinyin(s: str) -> str:
    """Каждое слово (и часть через дефис) с большой буквы, остальные буквы маленькие: LI JIANWEI -> Li Jianwei."""
    return re.sub(r"[A-Za-z]+", lambda m: m.group(0).capitalize(), s)


def lower_first(s: str) -> str:
    """Должность всегда с маленькой буквы: «Ронандаи экскватор» -> «ронандаи экскватор»."""
    return s[:1].lower() + s[1:]


def is_latin(s: str) -> bool:
    letters = [c for c in s if c.isalpha()]
    return bool(letters) and all(c.isascii() for c in letters)


# ------------------------------------------------------------- значения Excel
def norm(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v
    return str(v).strip()


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return dt.datetime.strptime(v, f).date()
            except ValueError:
                pass
    return None


def date_tj(d: dt.date) -> str:
    return f"{d.day} {MONTHS_TJ[d.month - 1]} соли {d.year}"


class Person:
    """Значения одного человека + правила производных полей."""

    def __init__(self, row: dict):
        self.row = {k: norm(v) for k, v in row.items()}
        self.notes = []                 # пометки для отчёта

    def raw(self, code):
        return self.row.get(code)       # None = колонка не выбрана, "" = ячейка пуста

    def value(self, code):
        v = self.raw(code)
        if code.startswith("ФИО"):
            return self._fio(code)
        if code in ("САНА", "САНА_СОЛИ", "СОЛ", "МОҲ", "РӮЗ"):
            return self._date(code, v)
        if isinstance(v, (dt.date, dt.datetime)):
            return to_date(v).strftime("%d.%m.%Y")
        if code.startswith("ВАЗИФА") and isinstance(v, str):
            return lower_first(v)
        return v

    def _date(self, code, v):
        d = to_date(v) if v not in (None, "") else None
        if d is None:                   # взять из соседней колонки с датой
            for other in ("САНА", "САНА_СОЛИ", "СОЛ"):
                o = self.row.get(other)
                if o and to_date(o):
                    d = to_date(o)
                    break
        if d is None:
            if code in ("СОЛ", "МОҲ", "РӮЗ"):
                s, m, r = (self.row.get(x) for x in ("СОЛ", "МОҲ", "РӮЗ"))
                return {"СОЛ": s, "МОҲ": m, "РӮЗ": r}[code] if self.raw(code) is not None else None
            if v is None and any(self.raw(c) is not None for c in ("САНА", "САНА_СОЛИ", "СОЛ")):
                return ""                # колонка с датой есть, ячейка пуста
            return v
        if code in ("САНА", "САНА_СОЛИ"):
            return v if (isinstance(v, str) and v and not to_date(v)) else date_tj(d)
        return str({"СОЛ": d.year, "МОҲ": d.month, "РӮЗ": d.day}[code])

    def _fio(self, code):
        cyr = lat = ""
        for name in ("ФИО", "ФИО_ТҶ", "ФИО_КИР"):
            if self.row.get(name):
                cyr = self.row[name]
                break
        lat = self.row.get("ФИО_ЛАТ", "") or ""
        if code not in ("ФИО", "ФИО_ЛАТ") and self.row.get(code):
            cyr, lat = self.row[code], ""
        lat = title_pinyin(lat)
        if not cyr and not lat and all(self.raw(c) is None for c in ("ФИО", "ФИО_ЛАТ", code)):
            return None
        if cyr and is_latin(cyr) and not lat:          # в «ФИО» лежит пиньин
            cyr, lat = "", title_pinyin(cyr)
        if lat and not cyr:
            cyr = latin_to_tj(lat)
            note = f"{code}: кириллица получена автотранслитом из «{lat}» — ПРОВЕРИТЬ"
            if note not in self.notes:
                self.notes.append(note)
        if not cyr and not lat:
            return ""
        if not FIO_BOTH_EVERYWHERE or not lat or cyr == lat:
            return cyr or lat
        return f"{cyr} ({lat})"

    def label(self):
        return self.row.get("ФИО") or self.row.get("ФИО_ЛАТ") or "(без ФИО)"


# ------------------------------------------------------------- замена в docx
def iter_paragraphs(doc):
    def walk(container):
        for p in container.paragraphs:
            yield p
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                     s.even_page_header, s.even_page_footer):
            if part is not None:
                yield from walk(part)


def replace_in_paragraph(p, resolve):
    """Заменяет коды {…}, даже если Word разбил их на несколько фрагментов (runs)."""
    pos = 0
    while True:
        runs = p.runs
        text = "".join(r.text for r in runs)
        m = PH.search(text, pos)
        if not m:
            return
        rep = resolve(m.group(1))
        start, end, acc = m.start(), m.end(), 0
        first = last = None
        for i, r in enumerate(runs):
            lo, hi = acc, acc + len(r.text)
            if first is None and lo <= start < hi:
                first, f_off = i, start - lo
            if lo < end <= hi:
                last, l_off = i, end - lo
                break
            acc = hi
        if first is None or last is None:
            return
        if first == last:
            t = runs[first].text
            runs[first].text = t[:f_off] + rep + t[l_off:]
        else:
            runs[first].text = runs[first].text[:f_off] + rep
            for k in range(first + 1, last):
                runs[k].text = ""
            runs[last].text = runs[last].text[l_off:]
        pos = start + len(rep)


def fill_template(path, person, out_path):
    doc = Document(path)
    stats = {"filled": [], "empty": [], "no_column": []}

    def resolve(code):
        v = person.value(code)
        if v is None:
            if code not in stats["no_column"]:
                stats["no_column"].append(code)
            return SPACES
        if v == "":
            if code not in stats["empty"]:
                stats["empty"].append(code)
            return SPACES
        if code not in stats["filled"]:
            stats["filled"].append(code)
        return str(v)

    for p in iter_paragraphs(doc):
        replace_in_paragraph(p, resolve)
    doc.save(out_path)
    return stats


# ----------------------------------------------------------------- диалоги
def ask_choice(title, options, multi=False):
    print(f"\n{title}")
    for i, o in enumerate(options, 1):
        print(f"  {i}. {o}")
    while True:
        s = input("Номер" + (" (можно несколько через запятую, «все» — все)" if multi else "") + ": ").strip().lower()
        try:
            if multi and s in ("все", "all", "*"):
                return list(range(len(options)))
            idx = [int(x) - 1 for x in s.replace(" ", "").split(",") if x]
            if idx and all(0 <= i < len(options) for i in idx) and (multi or len(idx) == 1):
                return idx if multi else idx[0]
        except ValueError:
            pass
        print("Не понял, повторите.")


def find_files(folder, exts):
    return sorted(f for f in os.listdir(folder)
                  if f.lower().endswith(exts) and not f.startswith("~$"))


def safe(s):
    return re.sub(r"[^\w\-. ()]+", "_", s).strip() or "bez_imeni"


def col_letter(i):
    return openpyxl.utils.get_column_letter(i + 1)


def short(v, n=18):
    s = "" if v is None else str(v).replace("\n", " ")
    return s if len(s) <= n else s[:n - 1] + "…"


def load_rows(xlsx, sheet):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]


def template_codes(paths):
    """Все коды {…} во всех выбранных шаблонах в порядке появления."""
    codes = []
    for t in paths:
        doc = Document(t)
        for p in iter_paragraphs(doc):
            for c in PH.findall(p.text):
                if c not in codes:
                    codes.append(c)
    return codes


FIRED = "САНАИ_ХОРИҶ"                   # служебный ключ: колонка с датой увольнения (фильтр, в шаблонах не нужен)


def _nk(x):
    """Для сравнения названий колонок: без регистра, пробелов и знаков; ҷ=ч, ё=е, ӣ=и."""
    x = str(x).lower().replace("ҷ", "ч").replace("ё", "е").replace("ӣ", "и").replace("ӯ", "у").replace("ҳ", "х")
    return re.sub(r"[\W_]+", "", x)


def auto_column(code, header):
    if code == FIRED:
        for i, h in enumerate(header):
            k = _nk(h)
            if k.startswith("санаихорич") or "увольн" in k or k in ("хорич", "хоричшудан"):
                return i
        return None
    for name in ALIASES.get(code, [code]):
        if name in header:
            return header.index(name)
    return None


def is_empty(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def choose_header_row(rows):
    print("\nПервые строки листа:")
    for i, r in enumerate(rows[:10], 1):
        print(f"  строка {i}: " + " | ".join(short(c) for c in r[:10]))
    default = next((i for i, r in enumerate(rows, 1) if any(c not in (None, "") for c in r)), 1)
    while True:
        s = input(f"В какой строке ШАПКА (заголовки)? [Enter = {default}]: ").strip()
        try:
            n = int(s) if s else default
            if 1 <= n <= len(rows):
                return n
        except ValueError:
            pass
        print("Не понял, повторите.")


def choose_mapping(codes, header, saved):
    """codes -> индекс колонки или None. Показывает таблицу, даёт править."""
    print("\nКолонки выбранной шапки:")
    for i, h in enumerate(header):
        print(f"  {col_letter(i):>3} (№{i + 1}): {h or '(пусто)'}")
    mapping = {}
    for c in codes:
        if c in saved and saved[c] is not None and saved[c] < len(header):
            mapping[c] = saved[c]
        else:
            mapping[c] = auto_column(c, header)
    while True:
        print("\nКлюч в шаблоне  →  колонка Excel")
        for n, c in enumerate(codes, 1):
            m = mapping[c]
            if c == FIRED:
                print(f"  {n:>2}. ⛔ Санаи хориҷ (фильтр: берём только тех, у кого колонка ПУСТАЯ)  →  "
                      + (f"{col_letter(m)}: {header[m] or '(пусто)'}" if m is not None else "— не выбрана (будут все работники)"))
                continue
            print(f"  {n:>2}. {{{c}}}  →  " + (f"{col_letter(m)}: {header[m] or '(пусто)'}" if m is not None else "— не выбрана (будут пробелы)"))
        s = input("\nEnter = всё верно; или номер ключа, который изменить: ").strip()
        if not s:
            return mapping
        try:
            code = codes[int(s) - 1]
        except (ValueError, IndexError):
            print("Не понял.")
            continue
        v = input(f"Для {'«Санаи хориҷ»' if code == FIRED else '{' + code + '}'} введите букву или номер колонки (A, B, 3 ...); 0 — не использовать: ").strip().upper()
        if v == "0":
            mapping[code] = None
        elif v.isdigit() and 1 <= int(v) <= len(header):
            mapping[code] = int(v) - 1
        elif v.isalpha() and 1 <= openpyxl.utils.column_index_from_string(v) <= len(header):
            mapping[code] = openpyxl.utils.column_index_from_string(v) - 1
        else:
            print("Такой колонки нет.")


def main():
    ap = argparse.ArgumentParser(description="Автозаполнение договоров и приказов из Excel")
    ap.add_argument("--xlsx"), ap.add_argument("--sheet")
    ap.add_argument("--templates", nargs="*")
    ap.add_argument("--out")
    ap.add_argument("--all", action="store_true", help="не фильтровать по «Санаи хориҷ» — взять всех, включая уволенных")
    ap.add_argument("--yes", action="store_true", help="не спрашивать про шапку и колонки (авто-подбор по названиям)")
    a = ap.parse_args()
    here = os.getcwd()

    xlsx = a.xlsx
    if not xlsx:
        xs = find_files(here, (".xlsx", ".xlsm"))
        if not xs:
            sys.exit("В текущей папке нет Excel-файла (.xlsx). Положите его рядом со скриптом или укажите --xlsx.")
        xlsx = xs[ask_choice("Какой Excel-файл использовать?", xs)]
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    sheet = a.sheet or wb.sheetnames[ask_choice("Какой ЛИСТ использовать?", wb.sheetnames)]
    wb.close()

    tpls = a.templates
    if not tpls:
        ds = [f for f in find_files(here, (".docx",)) if not f.startswith("zapolneno")]
        if not ds:
            sys.exit("В папке нет шаблонов .docx. Положите договор и приказ рядом или укажите --templates.")
        tpls = [ds[i] for i in ask_choice("Какие ШАБЛОНЫ заполнять?", ds, multi=True)]

    rows = load_rows(xlsx, sheet)
    if not rows:
        sys.exit("Лист пустой.")
    hrow = 1 if a.yes else choose_header_row(rows)
    if a.yes:
        hrow = next((i for i, r in enumerate(rows, 1) if any(c not in (None, "") for c in r)), 1)
    header = [str(h).strip() if h is not None else "" for h in rows[hrow - 1]]

    # ключи: из шаблонов + служебные, из которых считаются ФИО и даты
    codes = template_codes(tpls)
    for extra in ("ФИО", "ФИО_ЛАТ", "САНА", FIRED):
        if extra not in codes:
            codes.append(extra)

    map_file = f"mapping_{safe(os.path.splitext(os.path.basename(xlsx))[0])}_{safe(sheet)}.json"
    saved = {}
    if os.path.exists(map_file) and not a.yes:
        if input(f"\nНайден сохранённый выбор колонок ({map_file}). Использовать как основу? [Enter = да, n = нет]: ").strip().lower() != "n":
            with open(map_file, encoding="utf8") as f:
                saved = json.load(f)
    mapping = {c: auto_column(c, header) for c in codes} if a.yes else choose_mapping(codes, header, saved)
    with open(map_file, "w", encoding="utf8") as f:
        json.dump(mapping, f, ensure_ascii=False, indent=1)

    fi = None if a.all else mapping.get(FIRED)
    people, skipped = [], []                          # skipped — уволенные (есть дата в «Санаи хориҷ»)
    for r in rows[hrow:]:
        if all(c in (None, "") for c in r):
            continue
        row = {c: (r[i] if i is not None and i < len(r) else "") for c, i in mapping.items() if i is not None}
        p = Person(row)
        if fi is not None and not is_empty(r[fi] if fi < len(r) else None):
            d = to_date(r[fi]) if not isinstance(r[fi], str) else to_date(r[fi].strip())
            skipped.append((p.label(), d.strftime("%d.%m.%Y") if d else str(r[fi]).strip()))
            continue
        people.append(p)
    if not people:
        sys.exit("Нет ни одного действующего работника (у всех заполнена дата «Санаи хориҷ»)." if skipped
                 else "Под шапкой нет строк с данными.")
    filt = "выключен (--all)" if a.all else ("не задан: колонка «Санаи хориҷ» не найдена, взяты все" if fi is None
                                              else f"колонка {col_letter(fi)}: только с пустой датой")

    out = a.out or os.path.join(here, "zapolneno_" + dt.datetime.now().strftime("%Y%m%d_%H%M"))
    os.makedirs(out, exist_ok=True)

    report = []                                     # (человек, шаблон, файл, stats, notes)
    for n, person in enumerate(people, 1):
        folder = f"{n:02d}_{safe(str(person.label()))}"
        os.makedirs(os.path.join(out, folder), exist_ok=True)
        for t in tpls:
            name = f"{safe(os.path.splitext(os.path.basename(t))[0])}_{safe(str(person.label()))}.docx"
            person.notes = []
            st = fill_template(t, person, os.path.join(out, folder, name))
            report.append((person.label(), os.path.basename(t), f"{folder}/{name}", st, list(person.notes)))

    build_reports(out, xlsx, sheet, hrow, people, tpls, report, skipped, filt)
    warn = sum(1 for r in report if r[3]["empty"] or r[3]["no_column"] or r[4])
    print(f"\nГотово: людей {len(people)} (пропущено уволенных: {len(skipped)}), файлов {len(report)}, с пробелами/пометками: {warn}")
    if fi is None and not a.all:
        print("⚠ Колонка «Санаи хориҷ» не выбрана — взяты ВСЕ работники, включая уволенных.")
    print(f"Папка:  {out}")
    print(f"Отчёт:  {os.path.join(out, 'report.html')}  (и report.xlsx)")
    try:
        webbrowser.open("file://" + os.path.abspath(os.path.join(out, "report.html")))
    except Exception:
        pass


CSS = """
:root{--bg:#f4f6fa;--card:#fff;--ink:#1d2433;--mut:#6b7488;--line:#e3e7ef;--ok:#1f9d55;--okb:#e6f6ec;
--em:#b7791f;--emb:#fff4d9;--no:#c53030;--nob:#fde8e8;--wa:#6b46c1;--wab:#efe9fb;--ac:#2b5cd6}
@media (prefers-color-scheme:dark){:root{--bg:#141922;--card:#1c2330;--ink:#e8ecf4;--mut:#93a0b8;--line:#2b3446;
--ok:#4ad28a;--okb:#16352a;--em:#f0b955;--emb:#3b2f14;--no:#ff8a8a;--nob:#3d1c1c;--wa:#b79cf5;--wab:#2a2142;--ac:#7aa2ff}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 -apple-system,Segoe UI,Roboto,Arial,sans-serif}
.wrap{max-width:1080px;margin:0 auto;padding:28px 16px 60px}
h1{font-size:26px;margin:0 0 4px}.sub{color:var(--mut);margin-bottom:22px}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;margin-bottom:26px}
.kpi{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:14px 16px}
.kpi b{display:block;font-size:28px}.kpi span{color:var(--mut);font-size:13px}
.person{background:var(--card);border:1px solid var(--line);border-radius:14px;margin-bottom:18px;overflow:hidden}
.ph{display:flex;justify-content:space-between;align-items:center;gap:10px;padding:14px 18px;border-bottom:1px solid var(--line);flex-wrap:wrap}
.ph h2{margin:0;font-size:18px}.ph small{color:var(--mut)}
.badge{font-size:12px;font-weight:600;padding:3px 10px;border-radius:99px}
.b-ok{background:var(--okb);color:var(--ok)}.b-warn{background:var(--emb);color:var(--em)}
.doc{padding:14px 18px;border-bottom:1px dashed var(--line)}.doc:last-child{border-bottom:0}
.dt{font-weight:600;margin-bottom:8px}.dt span{font-weight:400;color:var(--mut);font-size:13px}
.row{display:flex;gap:8px;align-items:flex-start;margin:6px 0}
.lab{flex:0 0 150px;font-size:12px;font-weight:700;text-transform:uppercase;letter-spacing:.4px;padding-top:4px}
.l-ok{color:var(--ok)}.l-em{color:var(--em)}.l-no{color:var(--no)}.l-wa{color:var(--wa)}
.chips{display:flex;flex-wrap:wrap;gap:6px}
.chip{font:12px ui-monospace,Menlo,Consolas,monospace;padding:3px 9px;border-radius:8px}
.c-ok{background:var(--okb);color:var(--ok)}.c-em{background:var(--emb);color:var(--em)}
.c-no{background:var(--nob);color:var(--no)}.none{color:var(--mut);font-size:13px;padding-top:3px}
.note{background:var(--wab);color:var(--wa);padding:6px 10px;border-radius:8px;font-size:13px;margin:4px 0}
.leg{color:var(--mut);font-size:13px;margin:-8px 0 20px}
@media (max-width:640px){.lab{flex-basis:100%}.row{flex-direction:column;gap:2px}}
"""


def chips(items, cls):
    if not items:
        return '<span class="none">—</span>'
    return "".join(f'<span class="chip {cls}">{html.escape(x)}</span>' for x in items)


def build_reports(out, xlsx, sheet, hrow, people, tpls, report, skipped=(), filt=""):
    e = html.escape
    tot_f = sum(len(r[3]["filled"]) for r in report)
    tot_e = sum(len(r[3]["empty"]) for r in report)
    tot_n = sum(len(r[3]["no_column"]) for r in report)
    tot_w = sum(len(r[4]) for r in report)
    by_person = {}
    for r in report:
        by_person.setdefault(r[2].split("/")[0], []).append(r)
    parts = [f"<!doctype html><html lang='ru'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
             f"<title>Отчёт автозаполнения</title><style>{CSS}</style></head><body><div class='wrap'>",
             "<h1>Отчёт автозаполнения</h1>",
             f"<div class='sub'>{e(os.path.basename(xlsx))} · лист «{e(sheet)}» · шапка в строке {hrow} · "
             f"{dt.datetime.now():%d.%m.%Y %H:%M} · фильтр «Санаи хориҷ»: {e(filt)}</div>",
             "<div class='kpis'>"
             f"<div class='kpi'><b>{len(people)}</b><span>действующих работников</span></div>"
             f"<div class='kpi'><b style='color:var(--mut)'>{len(skipped)}</b><span>пропущено (есть «Санаи хориҷ»)</span></div>"
             f"<div class='kpi'><b>{len(report)}</b><span>документов ({len(tpls)} шаблона)</span></div>"
             f"<div class='kpi'><b style='color:var(--ok)'>{tot_f}</b><span>полей заполнено</span></div>"
             f"<div class='kpi'><b style='color:var(--em)'>{tot_e}</b><span>пусто в Excel (пробелы)</span></div>"
             f"<div class='kpi'><b style='color:var(--no)'>{tot_n}</b><span>колонка не выбрана (пробелы)</span></div>"
             f"<div class='kpi'><b style='color:var(--wa)'>{tot_w}</b><span>пометок — проверить</span></div></div>",
             "<div class='leg'>Зелёное — заполнено · жёлтое — ячейка в Excel пуста, в документе оставлены пробелы · "
             "красное — для этого ключа не выбрана колонка · фиолетовое — автоматические действия, которые стоит проверить</div>"]
    for folder, items in by_person.items():
        bad = any(i[3]["empty"] or i[3]["no_column"] or i[4] for i in items)
        parts.append(f"<div class='person'><div class='ph'><div><h2>{e(str(items[0][0]))}</h2><small>папка {e(folder)}</small></div>"
                     f"<span class='badge {'b-warn' if bad else 'b-ok'}'>{'есть что проверить' if bad else 'всё заполнено'}</span></div>")
        for who, tpl, fname, st, notes in items:
            parts.append(f"<div class='doc'><div class='dt'>{e(tpl)} <span>→ {e(fname)}</span></div>"
                         f"<div class='row'><div class='lab l-ok'>Заполнено · {len(st['filled'])}</div><div class='chips'>{chips(st['filled'], 'c-ok')}</div></div>"
                         f"<div class='row'><div class='lab l-em'>Пусто в Excel · {len(st['empty'])}</div><div class='chips'>{chips(st['empty'], 'c-em')}</div></div>"
                         f"<div class='row'><div class='lab l-no'>Нет колонки · {len(st['no_column'])}</div><div class='chips'>{chips(st['no_column'], 'c-no')}</div></div>")
            for n in notes:
                parts.append(f"<div class='note'>⚠ {e(n)}</div>")
            parts.append("</div>")
        parts.append("</div>")
    if skipped:
        parts.append("<div class='person'><div class='ph'><div><h2>Пропущены: уже уволены</h2><small>у них заполнена «Санаи хориҷ»</small></div>"
                     f"<span class='badge b-warn'>{len(skipped)}</span></div><div class='doc'><div class='chips'>"
                     + "".join(f"<span class='chip c-em'>{e(str(n))} — {e(str(d))}</span>" for n, d in skipped) + "</div></div></div>")
    parts.append("</div></body></html>")
    with open(os.path.join(out, "report.html"), "w", encoding="utf8") as f:
        f.write("".join(parts))

    from openpyxl.styles import Alignment, Font, PatternFill
    wo = openpyxl.Workbook()
    ws = wo.active
    ws.title = "Отчёт"
    ws.append(["Человек", "Шаблон", "Файл", "Заполнено", "Пусто в Excel", "Колонка не выбрана", "Пометки"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B5CD6")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for who, tpl, fname, st, notes in report:
        ws.append([who, tpl, fname, ", ".join(st["filled"]), ", ".join(st["empty"]),
                   ", ".join(st["no_column"]), "; ".join(notes)])
        r = ws.max_row
        for col, key, color in ((4, "filled", "E6F6EC"), (5, "empty", "FFF4D9"), (6, "no_column", "FDE8E8")):
            if st[key]:
                ws.cell(r, col).fill = PatternFill("solid", fgColor=color)
        if notes:
            ws.cell(r, 7).fill = PatternFill("solid", fgColor="EFE9FB")
    for col, w in zip("ABCDEFG", (24, 18, 42, 45, 38, 38, 48)):
        ws.column_dimensions[col].width = w
    if skipped:
        w2 = wo.create_sheet("Пропущены (уволены)")
        w2.append(["Человек", "Санаи хориҷ"])
        for c in w2[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="6B7488")
        for n, d in skipped:
            w2.append([n, d])
        w2.column_dimensions["A"].width = 36
        w2.column_dimensions["B"].width = 16
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wo.save(os.path.join(out, "report.xlsx"))


if __name__ == "__main__":
    main()
